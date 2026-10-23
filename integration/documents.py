"""Bounded real document extraction; filenames alone never count as extracted text."""
import codecs
import zipfile
from pathlib import Path
from state import UserError

MAX_FILE=64*1024*1024
MAX_CHARS=120000


def extract(path):
    path=Path(path)
    if not path.is_file():
        raise UserError('附件文件缺失')
    if path.stat().st_size>MAX_FILE:
        raise UserError('附件大于 64MB，未解析正文')
    head=path.open('rb')
    try:
        magic=head.read(16)
    finally:
        head.close()
    suffix=path.suffix.lower()
    if magic.startswith(b'%PDF'):
        from pypdf import PdfReader
        doc=PdfReader(path)
        if doc.is_encrypted and not doc.decrypt(''):
            raise UserError('PDF 受密码保护')
        chunks=[]; count=0
        for i,page in enumerate(doc.pages):
            text=page.extract_text() or ''
            chunks.append(f'\n[Page {i+1}]\n'+text)
            count+=len(chunks[-1])
            if count>MAX_CHARS:
                break
        text=''.join(chunks)
        if not any(c.strip() for c in text.replace('[Page','').splitlines() if not c.endswith(']')):
            raise UserError('扫描版 PDF 无可提取文字，需要 OCR')
        kind='pdf'
    elif magic.startswith(b'PK'):
        with zipfile.ZipFile(path) as z:
            infos=z.infolist()
            if len(infos)>10000 or sum(x.file_size for x in infos)>256*1024*1024 or any(x.flag_bits&1 for x in infos):
                raise UserError('Office 文件受保护或解压规模超过限制')
            names={x.filename for x in infos}
        chunks=[]; count=0
        def add(s):
            nonlocal count
            if count<=MAX_CHARS:
                chunks.append(s); count+=len(s)
        if 'word/document.xml' in names:
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            doc=Document(path)
            for node in doc.element.body.iterchildren():
                if node.tag==qn('w:p'):
                    p=Paragraph(node,doc)
                    add(f'[{p.style.name}] {p.text}')
                elif node.tag==qn('w:tbl'):
                    for row in Table(node,doc).rows:
                        add('\t'.join(c.text for c in row.cells))
                if count>MAX_CHARS: break
            kind='docx'
        elif 'xl/workbook.xml' in names:
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
            try:
                for sheet in book:
                    add(f'[Sheet {sheet.title}]')
                    for row in sheet.iter_rows():
                        values=[f'{c.coordinate}={c.value}' for c in row if c.value is not None]
                        if values: add('\t'.join(values))
                        if count>MAX_CHARS: break
                    if count>MAX_CHARS: break
            finally:
                book.close()
            kind='xlsx'
        elif 'ppt/presentation.xml' in names:
            from pptx import Presentation
            for i,slide in enumerate(Presentation(path).slides):
                add(f'[Slide {i+1}]')
                for shape in slide.shapes:
                    if shape.has_text_frame: add(shape.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            add('\t'.join(c.text for c in row.cells))
                if count>MAX_CHARS: break
            kind='pptx'
        else:
            raise UserError('不支持的 ZIP 文档格式')
        text='\n'.join(chunks)
    elif suffix in ('.txt','.md','.markdown'):
        raw=path.read_bytes()
        if raw.startswith((codecs.BOM_UTF16_LE,codecs.BOM_UTF16_BE)):
            text=raw.decode('utf-16')
        else:
            for enc in ('utf-8-sig','gb18030'):
                try:
                    text=raw.decode(enc); break
                except UnicodeError:
                    pass
            else:
                raise UserError('无法可靠识别文本编码')
        if '\x00' in text:
            raise UserError('文本含二进制数据或无 BOM 的 UTF-16，请转换编码')
        kind='markdown' if suffix!='.txt' else 'text'
    elif magic.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        raise UserError('旧 Office 格式或受密码保护的 Office 文档，暂不支持')
    else:
        raise UserError('实际文件类型暂不支持')
    if not text.strip():
        raise UserError('文档无可提取正文')
    return {'text':text[:MAX_CHARS],'type':kind,'truncated':len(text)>MAX_CHARS,
            'embedded_images':'未执行 OCR；图片引用保留在原附件'}
