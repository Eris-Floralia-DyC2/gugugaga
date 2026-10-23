import hashlib
import json
import os
import socket
import sys
import tempfile
import unittest
import wave
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from state import State,FileLock,atomic_json,UserError,digest
from cache import Cache,fingerprint
from wechat import parse_content
from wechat_reader import decrypt_page,mac_key,replay_wal
from fixtures import encrypted_db,encrypt_page,wal_file
from documents import extract
from analysis import input_message,acknowledge,prepare

def message(mid='one',content='original'):
    return {'id':mid,'timestamp':1,'content':content,'schema_version':1,'platform':'wechat','source_id':'s','account_id':'a',
            'chat_id':'chat','chat_name':'test','message_id':mid,'sender_id':'sender','sender_name':'test',
            'timezone':'Asia/Shanghai','message_type':1,'attachments':[],'reply_to':None,'parse_status':'ok','raw_reference':{},'truncated':False}


class Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=os.environ.get('WS_TEST_TMP'))
        self.root=Path(self.temp.name)
        self.s=State(self.root/'app/data')
        self.src=self.root/'original'; self.src.mkdir()
        self.key=b'K'*32; self.salt=b'S'*16
        self.db=self.src/'a.db'
        self.sql="CREATE TABLE t(value TEXT);INSERT INTO t VALUES('original');"

    def tearDown(self):
        self.temp.cleanup()

    def make(self):
        return encrypted_db(self.db,self.sql,self.key,self.salt)

    def test_cache_real_encrypted_roundtrip_reuse_source_readonly(self):
        self.make(); before=fingerprint(self.db)
        c=Cache(self.s)
        with c.open(self.src,'account-a','a.db',self.key) as (con,info):
            self.assertFalse(info['reused']);self.assertEqual(con.execute('SELECT value FROM t').fetchone()[0],'original')
            with self.assertRaises(Exception):con.execute("INSERT INTO t VALUES('forbidden')")
        with Cache(self.s).open(self.src,'account-a','a.db',self.key) as (_,info): self.assertTrue(info['reused'])
        self.assertEqual(before,fingerprint(self.db))

    def test_cache_corruption_rebuilt_and_account_isolated(self):
        self.make(); c=Cache(self.s)
        with c.open(self.src,'a','a.db',self.key) as (_,i): identity=i['id']
        (c.root/identity/'plain.db').write_bytes(b'corrupt')
        with c.open(self.src,'a','a.db',self.key) as (con,i): self.assertFalse(i['reused'])
        with c.open(self.src,'b','a.db',self.key) as (_,i): self.assertNotEqual(i['id'],identity)

    def test_same_mtime_content_change_rebuilds(self):
        self.make(); c=Cache(self.s)
        with c.open(self.src,'a','a.db',self.key): pass
        stamp=self.db.stat().st_mtime_ns
        # Alter source while keeping size/mtime; invalid MAC must fail rather than reuse.
        raw=bytearray(self.db.read_bytes());raw[130]^=1;self.db.write_bytes(raw)
        os.utime(self.db,ns=(stamp,stamp))
        with self.assertRaises(UserError):
            with c.open(self.src,'a','a.db',self.key):pass
        self.assertEqual(len(list(c.root.glob('*/plain.db'))),1)

    def test_page_authentication(self):
        self.make(); page=self.db.read_bytes()[:4096]
        with self.assertRaises(UserError):decrypt_page(page,1,b'J'*32,mac_key(b'J'*32,self.salt))
        corrupt=bytearray(page);corrupt[1024]^=1
        with self.assertRaises(UserError):decrypt_page(bytes(corrupt),1,self.key,mac_key(self.key,self.salt))

    def test_wal_update_visible_then_cache_reused(self):
        plain=self.make(); c=Cache(self.s)
        with c.open(self.src,'a','a.db',self.key):pass
        page=plain[4096:8192].replace(b'original',b'new_data')
        wal_file(Path(str(self.db)+'-wal'),[(2,2,encrypt_page(page,2,self.key,self.salt))])
        with c.open(self.src,'a','a.db',self.key) as (con,i):
            self.assertFalse(i['reused']);self.assertEqual(con.execute('SELECT value FROM t').fetchone()[0],'new_data')
        with c.open(self.src,'a','a.db',self.key) as (_,i):self.assertTrue(i['reused'])

    def test_wal_committed_only_preallocated_stale_tail(self):
        output=self.root/'out.db';output.write_bytes(b'Z'*4096+b'0'*4096)
        a=encrypt_page(b'A'*4016+bytes(80),2,self.key,self.salt)
        b=encrypt_page(b'B'*4016+bytes(80),2,self.key,self.salt)
        wal=self.root/'a.wal';wal_file(wal,[(2,2,a),(2,0,b)])
        with wal.open('ab') as f:f.write(bytes(5000))
        replay_wal(wal,output,self.key,mac_key(self.key,self.salt))
        self.assertEqual(output.read_bytes()[4096:],b'A'*4016+bytes(80))
        replay_wal(wal,output,self.key,mac_key(self.key,self.salt))
        self.assertEqual(output.read_bytes()[4096:],b'A'*4016+bytes(80))

    def test_lru_cannot_delete_leased_cache(self):
        self.make(); c=Cache(self.s)
        with c.open(self.src,'a','a.db',self.key):
            self.assertEqual(Cache(self.s).clean(all_unused=True)['removed'],0)
        self.assertEqual(c.clean(all_unused=True)['removed'],1)
        self.assertTrue(self.s.dbfile.exists())

    def test_snapshot_failure_retains_old_cache(self):
        self.make(); c=Cache(self.s)
        with c.open(self.src,'a','a.db',self.key) as (_,info): identity=info['id']
        output=c.root/identity/'plain.db';saved=output.read_bytes()
        raw=bytearray(self.db.read_bytes());raw[100]^=1;self.db.write_bytes(raw)
        with self.assertRaises(UserError):
            with c.open(self.src,'a','a.db',self.key):pass
        self.assertEqual(output.read_bytes(),saved)

    def test_disk_full_preserves_cache_and_recovers_on_retry(self):
        plain=self.make();c=Cache(self.s)
        with c.open(self.src,'a','a.db',self.key) as (_,i):identity=i['id']
        output=c.root/identity/'plain.db';saved=output.read_bytes()
        page=plain[4096:8192].replace(b'original',b'new_data')
        wal_file(Path(str(self.db)+'-wal'),[(2,2,encrypt_page(page,2,self.key,self.salt))])
        from types import SimpleNamespace
        with patch('cache.shutil.disk_usage',return_value=SimpleNamespace(free=0)):
            with self.assertRaises(UserError):
                with c.open(self.src,'a','a.db',self.key):pass
        self.assertEqual(output.read_bytes(),saved)
        with c.open(self.src,'a','a.db',self.key) as (con,_):self.assertEqual(con.execute('SELECT value FROM t').fetchone()[0],'new_data')

    def test_corrupt_cache_manifest_rebuilds_without_source_writes(self):
        self.make();c=Cache(self.s);before=fingerprint(self.db)
        with c.open(self.src,'a','a.db',self.key) as (_,i):identity=i['id']
        (c.root/identity/'meta.json').write_text('{broken',encoding='utf-8')
        with c.open(self.src,'a','a.db',self.key) as (con,i):self.assertFalse(i['reused'])
        self.assertEqual(before,fingerprint(self.db))
        orphan=c.root/identity/'build-interrupted';orphan.mkdir();(orphan/'partial').write_bytes(b'unfinished')
        with c.open(self.src,'a','a.db',self.key):pass
        self.assertFalse(orphan.exists())

    def test_messages_dedupe_and_atomic_cursor(self):
        m=message()
        self.s.save_messages('s','chat',[m],10)
        self.s.save_messages('s','chat',[m],10)
        self.assertEqual(len(self.s.messages()),1)
        with self.assertRaises(Exception): self.s.save_messages('s','chat',[{'broken':True}],20)
        with self.s.db() as con:self.assertEqual(con.execute('SELECT until FROM cursors').fetchone()[0],10)

    def test_same_server_message_in_two_shards_does_not_reopen_analysis(self):
        from wechat import WeChat
        root=self.src/'db_storage';chat='chat@chatroom'
        table='Msg_'+hashlib.md5(chat.encode()).hexdigest()
        sql=f"CREATE TABLE {table}(local_id INTEGER,local_type INTEGER,create_time INTEGER,message_content TEXT,server_id INTEGER); INSERT INTO {table} VALUES(1,1,1,'same message',77);"
        files={'contact/contact.db':"CREATE TABLE contact(username TEXT,nick_name TEXT,remark TEXT); INSERT INTO contact VALUES('chat@chatroom','test','');",'message/message_0.db':sql,'message/message_1.db':sql}
        for rel,body in files.items():encrypted_db(root/rel,body,self.key,self.salt)
        source={'id':'s','platform':'wechat','account_id':root.parent.name,'db_dir':str(root),'groups':[chat],'enabled':True}
        self.s.secret('wechat-s',{'db_dir':str(root),'keys':{rel:self.key.hex() for rel in files}},write=True)
        reader=WeChat(self.s,source)
        rows=list(reader.iter_messages(chat,0,2));self.assertEqual(len(rows),1)
        self.s.save_messages('s',chat,rows,2)
        with self.s.db() as con:con.execute('UPDATE messages SET analyzed=revision')
        self.s.save_messages('s',chat,list(reader.iter_messages(chat,0,2)),2)
        self.assertFalse(self.s.messages(pending=True))

    def test_xml_reply_and_safe_malformed(self):
        raw='<msg><appmsg><title>交作业</title><type>57</type><refermsg><svrid>123</svrid><content>周五</content></refermsg></appmsg></msg>'
        text,ref,_,status=parse_content(raw,49)
        self.assertIn('周五',text);self.assertEqual(ref['message_id'],'123');self.assertEqual(status,'ok')
        self.assertEqual(parse_content('<msg><broken>',49)[3],'xml_failed')

    def test_remote_scope_excludes_document_and_transcript(self):
        m={'id':'m','chat_id':'c','chat_name':'n','timestamp':1,'sender_name':'x','message_type':1,'content':'visible',
           'attachments':[{'kind':'document','name':'private','extracted_text':'secret'}]}
        item=input_message(m,{'text':True,'documents':False,'transcriptions':False})
        self.assertEqual(item['text'],'visible')
        m['attachments']=[{'kind':'image','extracted_text':'uncertain OCR','vision_uncertainty':'文字模糊'}]
        item=input_message(m,{'text':True,'documents':False,'transcriptions':False,'images':True})
        self.assertTrue(item['truncated']);self.assertIn('文字模糊',item['text'])
        m.update(message_type=34,content='<voicemsg aeskey="private-transport-token"/>',attachments=[])
        item=input_message(m,{'text':True,'documents':False,'transcriptions':False,'images':False})
        self.assertNotIn('private-transport-token',item['text'])

    def test_attachment_results_survive_recollection_without_reextracting(self):
        from attachments import Attachments
        from types import SimpleNamespace
        adapter=SimpleNamespace(root=self.src/'db_storage',source={'id':'s'})
        processor=Attachments(self.s,adapter)
        file=self.src/'body.txt';file.write_text('a real document body',encoding='utf-8')
        m=message();m['attachments']=[{'kind':'document','name':'body.txt','stages':{'identified':'ok'}}]
        with patch.object(processor,'locate',return_value=file):processor.process(m)
        self.s.save_messages('s','chat',[m],2)
        fresh=message();fresh['attachments']=[{'kind':'document','name':'body.txt','stages':{'identified':'ok'}}]
        with patch.object(processor,'locate',return_value=file),patch('attachments.extract',side_effect=AssertionError('should reuse')):
            processor.process(fresh)
        self.assertEqual(m,fresh)
        self.assertEqual(fresh['attachments'][0]['extracted_text'],'a real document body')

    def test_outbox_ack_does_not_mark_newer_message_analyzed(self):
        old=message('m','old');new={**old,'content':'changed'}
        self.s.save_messages('s','chat',[old],1)
        payload={'tasks':[],'revisions':{'m':digest(old)}}
        with self.s.db() as con:con.execute('INSERT INTO analyses VALUES(?,?,?,?)',('a','ready',json.dumps(payload),1))
        self.s.save_messages('s','chat',[new],2)
        acknowledge(self.s,'a',[])
        self.assertEqual(len(self.s.messages(pending=True)),1)

    def test_bad_llm_response_and_network_failure_do_not_lose_messages(self):
        c=self.s.config();c['sources']=[{'id':'s','platform':'wechat','account_id':'a','account_name':'测试','db_dir':'','enabled':True}];atomic_json(self.s.config_file,c)
        atomic_json(self.s.root/'config/summary.json',{'version':1,'selected':[{'source_id':'s','chat_id':'chat@chatroom'}],'since':0,'until':10})
        m=message();m['chat_id']='chat@chatroom';self.s.save_messages('s',m['chat_id'],[m],10)
        with patch('analysis.request',side_effect=UserError('network unavailable')):
            with self.assertRaises(UserError):prepare(self.s,{'key':'test'},[])
        self.assertEqual(len(self.s.messages(pending=True)),1)
        with patch('analysis.request',return_value={'unexpected':'invalid'}):
            with self.assertRaises(UserError):prepare(self.s,{'key':'test'},[])
        self.assertEqual(len(self.s.messages(pending=True)),1)

    def test_clear_assignment_missing_deadline_is_valid_without_invented_date(self):
        from proposal_rules import validate_proposals
        m={'id':'m','group_id':'chat@chatroom','group_name':'课程','timestamp':1,'sender':'老师','text':'下次物理课交作业'}
        task={'title':'交作业','description':'下次课交','plan_date':'2026-10-20','due_date':None,'priority':'normal',
              'source_ids':['m'],'evidence':[{'message_id':'m','quote':'下次物理课交作业'}],'applies_to_me':True,
              'uncertainty':'没有具体截止日期','confidence':.8,'existing_task_id':None}
        accepted,pending,_=validate_proposals({'tasks':[task],'review_notes':[]},[m],[])
        self.assertEqual(len(accepted),1);self.assertFalse(pending);self.assertIsNone(accepted[0]['dueDate'])
        task['uncertainty']='';task['due_date']='2099-01-01'
        accepted,_,_=validate_proposals({'tasks':[task],'review_notes':[]},[m],[])
        self.assertIsNone(accepted[0]['dueDate'])
        task['evidence'][0]['quote']='编造的要求'
        with self.assertRaises(UserError):validate_proposals({'tasks':[task],'review_notes':[]},[m],[])

    def test_mcp_denies_outside_allowlist(self):
        from mcp_server import authorize,settings
        settings(self.s)
        with self.assertRaises(UserError):authorize(self.s,'s','c',0,10,10)

    def test_dpapi_roundtrip(self):
        self.s.secret('test',{'a':'private'},write=True)
        self.assertEqual(self.s.secret('test'),{'a':'private'})
        self.assertNotIn(b'private',(self.s.root/'secrets/test.dpapi').read_bytes())

    def test_six_document_formats_actual_body(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        from reportlab.pdfgen.canvas import Canvas
        p=self.root/'a.txt';p.write_text('正文 TEST 内容',encoding='utf-16')
        self.assertIn('TEST',extract(p)['text'])
        p=self.root/'a.md';p.write_text('# TEST\n正文',encoding='utf-8')
        self.assertIn('# TEST',extract(p)['text'])
        p=self.root/'a.pdf';can=Canvas(str(p));can.drawString(40,700,'TEST pdf body');can.save()
        self.assertIn('TEST pdf body',extract(p)['text'])
        p=self.root/'a.docx';doc=Document();doc.add_heading('TEST',1);doc.add_table(1,1).cell(0,0).text='table body';doc.save(p)
        self.assertIn('table body',extract(p)['text'])
        p=self.root/'a.xlsx';book=Workbook();book.active['B2']='TEST cell';book.save(p);book.close()
        self.assertIn('B2=TEST cell',extract(p)['text'])
        p=self.root/'a.pptx';slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1)).text='TEST slide';slides.save(p)
        self.assertIn('TEST slide',extract(p)['text'])

    def test_image_xor_and_v1_v2_real_bytes(self):
        from PIL import Image
        from Crypto.Cipher import AES
        from Crypto.Util.Padding import pad
        from decode_image import decrypt_dat_file
        p=self.root/'original.png';Image.new('RGB',(20,20),(20,40,60)).save(p);raw=p.read_bytes()
        dat=self.root/'image.dat';dat.write_bytes(bytes(b^0x71 for b in raw))
        out=self.root/'out.png';self.assertTrue(decrypt_dat_file(str(dat),str(out))[0]);self.assertEqual(raw,out.read_bytes())
        import struct
        for version,key in [('V1',b'cfcd208495d565ef'),('V2',b'K'*16)]:
            aes_size=32;xor_size=16
            dat.write_bytes(b'\x07\x08'+version.encode()+b'\x08\x07'+struct.pack('<II',aes_size,xor_size)+b'\0'+AES.new(key,AES.MODE_ECB).encrypt(pad(raw[:aes_size],16))+raw[aes_size:-xor_size]+bytes(b^0x88 for b in raw[-xor_size:]))
            self.assertTrue(decrypt_dat_file(str(dat),str(out),key)[0]);self.assertEqual(raw,out.read_bytes())

    def test_voice_is_unprocessed_never_located_or_transcribed(self):
        from attachments import Attachments
        from types import SimpleNamespace
        reader=SimpleNamespace(root=self.src/'db_storage',source={'id':'s'},voice=lambda _:self.fail('voice bytes must not be read'))
        m=message();m['message_type']=34;m['attachments']=[{'kind':'audio','stages':{}}]
        processor=Attachments(self.s,reader)
        with patch.object(processor,'locate',side_effect=AssertionError('must not locate')):processor.process(m)
        self.assertEqual(m['attachments'][0]['stages']['processed'],'unprocessed')
        self.assertEqual(input_message(m,{'text':True,'documents':True,'images':True})['text'],'')


if __name__=='__main__':unittest.main()
